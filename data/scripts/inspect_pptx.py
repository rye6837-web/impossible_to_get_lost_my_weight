import sys
from pptx import Presentation

try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

prs = Presentation(r'presentation\AI_Diet_Coach_Presentation_Interim.pptx')
print(f"Total Slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for j, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            txt = shape.text.strip().replace('\n', ' // ')
            if txt:
                print(f"  [Shape {j}] {txt[:120]}")
        elif shape.has_table:
            print(f"  [Shape {j} - Table ({len(shape.table.rows)}x{len(shape.table.columns)})]")
        else:
            print(f"  [Shape {j} - {shape.shape_type}]")
