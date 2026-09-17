import sys
from pptx import Presentation

try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

prs = Presentation(r'presentation\AI_Diet_Coach_Presentation_Interim.pptx')

def dump_slide(idx):
    slide = prs.slides[idx - 1]
    print(f"\n==================== SLIDE {idx} ====================")
    for j, shape in enumerate(slide.shapes):
        print(f"Shape {j} [{shape.name} | type={shape.shape_type} | pos=({shape.left.inches:.2f}, {shape.top.inches:.2f}) size=({shape.width.inches:.2f}, {shape.height.inches:.2f})]")
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                p_text = p.text.strip()
                if p_text:
                    print(f"    P{p_idx} (font={p.font.name}, sz={p.font.size.pt if p.font.size else 'None'}): {p_text[:100]}")

for s in [2, 4, 5, 6, 10, 13, 15, 16]:
    dump_slide(s)
